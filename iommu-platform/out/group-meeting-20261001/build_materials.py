#!/usr/bin/env python3
# -*- coding: utf-8 -*-
'''Group-meeting material builder.  Local only, no network, no input mutation.'''
from __future__ import annotations
import argparse, csv, hashlib, json, math, os, shutil, statistics, subprocess, sys, tempfile, zipfile
from pathlib import Path
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
NL = chr(10)
REPORT_DATE = '2026-10-01'
CUTOFF_BOARD = '2026-09-20'
CUTOFF_ORGANIZED = '2026-09-30'
TARGET_SECONDS = 540
MIN_SECONDS = 480
MAX_SECONDS = 600
EXPECTED_SLIDES = 8
CLOCK_HZ = 50000000.0
CYCLE_NS = 20.0
MIB = float(1 << 20)
SCOPE_TOTAL = 60
ONLINE = 24
OFFLINE = 1
PENDING = 35
E3_RAW_REL = 'out/mac-e3-controlled-20260918-164635-run-rk48bhte/all-original-pairs.json'
E3_REVIEW_REL = 'out/e3-controlled-review-20260918/review-summary.json'
P4M_TRIALS_REL = 'out/mac-p4m-formal60-20260920-215406-run-cicggil4/trial_summary.json'
P4M_STATS_REL = 'out/p4m-formal60-recovery-20260920/statistics.json'
NUMERIC_NAME = 'numeric-verification.json'
SLIDES_CANDIDATES = ['out/slides.json', 'out/group-meeting-20261001/slides.json']
DARK_BLUE = RGBColor(0x0B, 0x2E, 0x59)
CYAN = RGBColor(0x00, 0x9A, 0xA6)
ORANGE = RGBColor(0xE6, 0x7E, 0x22)
GRAY = RGBColor(0x55, 0x5F, 0x6D)
LIGHT = RGBColor(0xF2, 0xF6, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x2E, 0x7D, 0x32)
RED = RGBColor(0xB0, 0x2A, 0x2A)
CJK_FONT_NAME = None
FONT_PATH = None
PPT_CJK_FONT_NAME = 'Noto Sans CJK SC'
def sha256_file(path):
    h = hashlib.sha256()
    with Path(path).open('rb') as f:
        for chunk in iter(lambda: f.read(1 << 20), b''):
            h.update(chunk)
    return h.hexdigest()
def load_json(path):
    with Path(path).open('r', encoding='utf-8') as f:
        return json.load(f)
def write_json(path, obj):
    with Path(path).open('w', encoding='utf-8') as f:
        json.dump(obj, f, ensure_ascii=False, indent=2, sort_keys=True)
        f.write(NL)
def write_csv(path, header, rows):
    with Path(path).open('w', newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)
def rel(path, root):
    try:
        return str(Path(path).resolve().relative_to(Path(root).resolve()))
    except Exception:
        return str(path)
def find_cjk_font():
    global CJK_FONT_NAME, FONT_PATH
    cands = []
    for name in ['Noto Sans CJK SC', 'Noto Sans CJK SC Regular', 'Noto Sans CJK SC Bold']:
        try:
            p = subprocess.check_output(['fc-match', '-f', '%{file}', name], text=True, stderr=subprocess.DEVNULL).strip()
            if p and Path(p).exists():
                cands.append(p)
        except Exception:
            pass
    if not cands:
        raise RuntimeError('Noto Sans CJK SC font not found by fc-match')
    FONT_PATH = cands[0]
    font_manager.fontManager.addfont(FONT_PATH)
    CJK_FONT_NAME = font_manager.FontProperties(fname=FONT_PATH).get_name()
    matplotlib.rcParams['font.sans-serif'] = [CJK_FONT_NAME]
    matplotlib.rcParams['axes.unicode_minus'] = False
    return CJK_FONT_NAME, FONT_PATH
def fprop(size=None, bold=False):
    if CJK_FONT_NAME is None:
        find_cjk_font()
    return font_manager.FontProperties(fname=FONT_PATH, size=size, weight='bold' if bold else 'normal')
def set_run_font(run, size, bold=False, color=DARK_BLUE):
    name = PPT_CJK_FONT_NAME
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ['a:ea', 'a:cs']:
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set('typeface', name)
def add_text(slide, left, top, width, height, lines, size=20, bold=False, color=DARK_BLUE, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, spacing=1.02):
    if isinstance(lines, str):
        lines = [lines]
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        r = p.add_run()
        r.text = str(line)
        set_run_font(r, size, bold=(bold and i == 0), color=color)
    return tb
def add_title(slide, text, size=30):
    if len(text) > 24:
        size = 26
    return add_text(slide, Inches(0.45), Inches(0.25), Inches(12.45), Inches(0.85), text, size=size, bold=True, color=DARK_BLUE, spacing=0.95)
def add_takeaway(slide, text):
    return add_text(slide, Inches(0.55), Inches(1.12), Inches(12.25), Inches(0.78), text, size=22, bold=True, color=DARK_BLUE, spacing=0.98)
def add_footer(slide, num, sources):
    def short_source(s):
        s = str(s).strip()
        base = s
        line = ''
        if ':' in s:
            head, tail = s.rsplit(':', 1)
            if tail.strip().isdigit():
                base = head
                line = tail.strip()
        name = Path(base).name
        if len(name) > 24:
            p = Path(name)
            name = p.stem[:20] + '…' + p.suffix
        return name + (':' + line if line else '')
    src = short_source(sources[0]) if sources else '见备注'
    txt = REPORT_DATE + '｜板端数据截至' + CUTOFF_BOARD[5:] + '｜整理' + CUTOFF_ORGANIZED[5:] + '｜来源：' + src
    if num:
        txt += '｜第' + str(num) + '页'
    return add_text(slide, Inches(0.45), Inches(7.05), Inches(12.45), Inches(0.32), txt, size=11, color=GRAY)
def add_notes(slide, sd):
    lines = []
    lines.append('标题：' + str(sd.get('title', '')))
    lines.append('口语稿：')
    notes = sd.get('notes', [])
    if isinstance(notes, str):
        notes = [notes]
    lines.extend([str(x) for x in notes])
    lines.append('')
    lines.append('来源：')
    lines.extend([str(x) for x in sd.get('sources', [])])
    lines.append('报告日期：' + REPORT_DATE)
    lines.append('资料截止：最新板端 ' + CUTOFF_BOARD + '，资料整理 ' + CUTOFF_ORGANIZED)
    lines.append('说明：本页未声称新板测试；口语稿仅存备注。')
    slide.notes_slide.notes_text_frame.text = NL.join(lines)
def e3_w(r, i):
    arr = r['e3_words']
    if i < 0 or i >= len(arr):
        raise AssertionError('e3_words index out of range ' + str(i))
    return arr[i]
def e3_one(raw_obj, index0, orig1, main_case, pair):
    assert raw_obj.get('success') is True, 'E3 raw success false at pair ' + str(orig1)
    inp = raw_obj.get('input', {})
    assert inp.get('chunk_bytes') == 64, 'E3 not 64B at pair ' + str(orig1)
    assert inp.get('transfers') == 1, 'E3 transfers not 1 at pair ' + str(orig1)
    assert inp.get('prep') == 'cold', 'E3 not cold at pair ' + str(orig1)
    assert inp.get('pair') == 'complementary', 'E3 not complementary at pair ' + str(orig1)
    rounds = {int(r['delay']): r for r in raw_obj.get('rounds', [])}
    assert set(rounds.keys()) == {0, 128}, 'E3 delays not 0/128 at pair ' + str(orig1)
    vals = {}
    for d in (0, 128):
        r = rounds[d]
        a = e3_w(r, 41) - e3_w(r, 4)
        v = e3_w(r, 70) - e3_w(r, 4)
        own = e3_w(r, 73) - e3_w(r, 71)
        assert own == 4, 'V own per-round not 4 at pair ' + str(orig1) + ' delay ' + str(d) + ' got ' + str(own)
        vals[d] = {'a_remaining': a, 'v_remaining': v, 'v_own': own}
    delta = {'a_remaining': vals[128]['a_remaining'] - vals[0]['a_remaining'], 'v_remaining': vals[128]['v_remaining'] - vals[0]['v_remaining'], 'v_own': vals[128]['v_own'] - vals[0]['v_own']}
    assert delta['v_own'] == 0, 'V own delta not zero at pair ' + str(orig1)
    return {'pair_index0': index0, 'original_pair_1based': orig1, 'main_case': main_case, 'values': vals, 'delta': delta}
def build_e3(root, out_dir, numeric):
    raw = load_json(root / E3_RAW_REL)
    rev = load_json(root / E3_REVIEW_REL)
    assert isinstance(raw, list) and len(raw) >= 11, 'E3 raw list too short'
    specs = [(4, 5, 2), (6, 7, 4), (8, 9, 6), (10, 11, 8)]
    recs = []
    for idx, orig1, main_case in specs:
        rec = e3_one(raw[idx], idx, orig1, main_case, raw[idx].get('input', {}))
        recs.append(rec)
        rp = [p for p in rev.get('pairs', []) if p.get('phase') == 'main' and p.get('case') == main_case and p.get('role') == '1_cold_complementary']
        assert len(rp) == 1, 'review pair not found main case ' + str(main_case)
        b = rp[0].get('baseline', {})
        for k in ['A_remaining', 'V_remaining', 'V_own']:
            expk = 'a_remaining' if k == 'A_remaining' else ('v_remaining' if k == 'V_remaining' else 'v_own')
            assert b.get(k) == rec['delta'][expk], 'review baseline mismatch ' + str(main_case) + ' ' + k
        np = [x for x in numeric.get('e3_pairs', []) if x.get('original_pair_1based') == orig1 and x.get('main_case') == main_case]
        assert len(np) == 1, 'numeric e3 pair not found ' + str(orig1)
        for d in (0, 128):
            nv = np[0]['values'][str(d)]
            assert int(nv['a_remaining']) == rec['values'][d]['a_remaining'], 'numeric A mismatch pair ' + str(orig1) + ' d' + str(d)
            assert int(nv['v_remaining']) == rec['values'][d]['v_remaining'], 'numeric V mismatch pair ' + str(orig1) + ' d' + str(d)
            assert int(nv['v_own']) == rec['values'][d]['v_own'], 'numeric own mismatch pair ' + str(orig1) + ' d' + str(d)
        nd = np[0]['delta']
        assert nd['a_remaining'] == rec['delta']['a_remaining'], 'numeric delta A mismatch pair ' + str(orig1)
        assert nd['v_remaining'] == rec['delta']['v_remaining'], 'numeric delta V mismatch pair ' + str(orig1)
        assert nd['v_own'] == rec['delta']['v_own'], 'numeric delta own mismatch pair ' + str(orig1)
    assert [r['delta']['a_remaining'] for r in recs] == [127, 129, 129, 130], 'A deltas not 127/129/129/130'
    assert [r['delta']['v_remaining'] for r in recs] == [127, 129, 129, 130], 'V deltas not 127/129/129/130'
    figs = out_dir / 'figures'
    figs.mkdir(parents=True, exist_ok=True)
    rows = []
    for r in recs:
        rows.append([r['pair_index0'], r['original_pair_1based'], r['main_case'], r['values'][0]['a_remaining'], r['values'][0]['v_remaining'], r['values'][0]['v_own'], r['values'][128]['a_remaining'], r['values'][128]['v_remaining'], r['values'][128]['v_own'], r['delta']['a_remaining'], r['delta']['v_remaining'], r['delta']['v_own'], 'rounds delay0/128 e3_words w41-w4/w70-w4/w73-w71; raw=' + E3_RAW_REL])
    write_csv(figs / 'e3-pairs.csv', ['pair_index0', 'original_pair_1based', 'main_case', 'd0_a_remaining', 'd0_v_remaining', 'd0_v_own', 'd128_a_remaining', 'd128_v_remaining', 'd128_v_own', 'delta_a_remaining', 'delta_v_remaining', 'delta_v_own', 'source'], rows)
    x = list(range(len(recs)))
    a = [r['delta']['a_remaining'] for r in recs]
    v = [r['delta']['v_remaining'] for r in recs]
    labels = ['对照1', '对照2', '对照3', '对照4']
    fig, ax = plt.subplots(figsize=(8.6, 4.8), dpi=200)
    w = 0.34
    b1 = ax.bar([i - w/2 for i in x], a, width=w, color='#E67E22', label='A剩余服务增量')
    b2 = ax.bar([i + w/2 for i in x], v, width=w, color='#009AA6', label='V剩余等待增量')
    ax.axhline(128, color='#555F6D', linestyle='--', linewidth=1.0)
    ax.text(0.02, 0.78, '人工插入128周期参考（非自然常数）', transform=ax.transAxes, color='#555F6D', fontproperties=fprop(9), va='center')
    for bars in (b1, b2):
        for rect in bars:
            h = rect.get_height()
            ax.annotate(str(int(h)), (rect.get_x()+rect.get_width()/2, h), textcoords='offset points', xytext=(0,3), ha='center', fontproperties=fprop(9))
    ax.set_xticks(x)
    ax.set_xticklabels(labels, fontproperties=fprop(11))
    ax.set_ylabel('额外等待/服务时间（周期）', fontproperties=fprop(11))
    ax.set_title('让A多等，V也多等', fontproperties=fprop(14, True))
    ax.set_ylim(0, max(max(a), max(v), 128) * 1.35)
    ax.legend(prop=fprop(10), frameon=False, loc='lower right')
    ax.grid(axis='y', color='#DDDDDD', linewidth=0.7)
    caption = NL.join([
        '单位：时钟周期。四组依次对应 maincase2/4/6/8（描述性，64 B 诊断）。',
        '128 周期为受控干预人为插入，不是自然常数；只延迟 A 的一次页表响应。',
        'A 剩余服务增量与 V 剩余等待增量同步出现；V 自身处理 4 周期不变。',
        '范围受限：本实现、本配置、本到达方式；不是 V 完整排队或总耗时增量。',
    ])
    fig.text(0.02, 0.02, caption, fontproperties=fprop(8), color='#555F6D', linespacing=1.25, va='bottom')
    fig.tight_layout(rect=[0,0.20,1,1])
    fig.savefig(figs / '机制干预证据.png', dpi=200)
    fig.savefig(figs / '机制干预证据.pdf')
    plt.close(fig)
    return {'records': recs, 'figure_png': str(figs / '机制干预证据.png'), 'figure_pdf': str(figs / '机制干预证据.pdf'), 'csv': str(figs / 'e3-pairs.csv')}
def p4m_rate(byte_count, cycles):
    return float(byte_count) / (float(cycles) / CLOCK_HZ) / MIB
def build_p4m(root, out_dir, stats):
    trials = load_json(root / P4M_TRIALS_REL)
    assert isinstance(trials, list) and len(trials) >= 25, 'P4M trial list too short'
    first24 = trials[:24]
    assert len(first24) == 24, 'P4M first24 length wrong'
    conds = ['gap0', 'gap1000', 'gap2000']
    recs = []
    blocks = {}
    counts = {c: 0 for c in conds}
    for t in first24:
        assert t.get('completed') is True, 'P4M incomplete in first24 ' + str(t.get('trial_id'))
        assert t.get('capture') == 'complete', 'P4M capture not complete ' + str(t.get('trial_id'))
        assert t.get('case_exit') == 0 and t.get('case_end') == 0, 'P4M case exit/end not zero ' + str(t.get('trial_id'))
        c = t.get('condition')
        assert c in conds, 'P4M bad condition ' + str(c)
        counts[c] += 1
        blocks.setdefault(int(t.get('block')), set()).add(c)
        v = p4m_rate(t['victim_completed_bytes'], t['v_cycles'])
        a = p4m_rate(t['a_window_bytes'], t['v_cycles'])
        recs.append({'index': t.get('index'), 'trial_id': t.get('trial_id'), 'block': t.get('block'), 'gap': t.get('gap'), 'condition': c, 'sequence': t.get('sequence'), 'v_cycles': t.get('v_cycles'), 'window_ns': t.get('window_ns'), 'victim_completed_bytes': t.get('victim_completed_bytes'), 'a_window_bytes': t.get('a_window_bytes'), 'a_window_jobs': t.get('a_window_jobs'), 'v_mib_s': v, 'a_mib_s': a, 'source': P4M_TRIALS_REL})
    for c in conds:
        assert counts[c] == 8, 'P4M condition count not 8 ' + c
    for b in range(1, 9):
        assert blocks.get(b) == set(conds), 'P4M block not balanced ' + str(b)
    figs = out_dir / 'figures'
    figs.mkdir(parents=True, exist_ok=True)
    write_csv(figs / 'p4m-trials-first24.csv', ['index', 'trial_id', 'block', 'gap', 'condition', 'sequence', 'v_cycles', 'window_ns', 'victim_completed_bytes', 'a_window_bytes', 'a_window_jobs', 'v_mib_s', 'a_mib_s', 'source'], [[r[k] for k in ['index', 'trial_id', 'block', 'gap', 'condition', 'sequence', 'v_cycles', 'window_ns', 'victim_completed_bytes', 'a_window_bytes', 'a_window_jobs', 'v_mib_s', 'a_mib_s', 'source']] for r in recs])
    sums = []
    for c in conds:
        xs = [r for r in recs if r['condition'] == c]
        vm = [r['v_mib_s'] for r in xs]
        am = [r['a_mib_s'] for r in xs]
        vc = [r['v_cycles'] for r in xs]
        wn = [r['window_ns'] for r in xs]
        st = stats['primary_per_condition'][c]
        for key, vals in [('v_mib_s', vm), ('a_window_mib_s', am)]:
            e = st[key]
            got = [min(vals), statistics.median(vals), max(vals)]
            exp = [e['min'], e['median'], e['max']]
            for g, z in zip(got, exp):
                assert abs(float(g) - float(z)) <= 1e-10, 'P4M stats mismatch ' + c + ' ' + key + ' ' + str(g) + ' vs ' + str(z)
        sums.append({'condition': c, 'n': len(xs), 'v_mib_s_min': min(vm), 'v_mib_s_median': statistics.median(vm), 'v_mib_s_max': max(vm), 'a_mib_s_min': min(am), 'a_mib_s_median': statistics.median(am), 'a_mib_s_max': max(am), 'v_cycles_min': min(vc), 'v_cycles_median': statistics.median(vc), 'v_cycles_max': max(vc), 'window_ns_min': min(wn), 'window_ns_median': statistics.median(wn), 'window_ns_max': max(wn)})
    write_csv(figs / 'p4m-condition-summary.csv', ['condition', 'n', 'v_mib_s_min', 'v_mib_s_median', 'v_mib_s_max', 'a_mib_s_min', 'a_mib_s_median', 'a_mib_s_max', 'v_cycles_min', 'v_cycles_median', 'v_cycles_max', 'window_ns_min', 'window_ns_median', 'window_ns_max'], [[s[k] for k in ['condition', 'n', 'v_mib_s_min', 'v_mib_s_median', 'v_mib_s_max', 'a_mib_s_min', 'a_mib_s_median', 'a_mib_s_max', 'v_cycles_min', 'v_cycles_median', 'v_cycles_max', 'window_ns_min', 'window_ns_median', 'window_ns_max']] for s in sums])
    cmap = {'gap0': '#009AA6', 'gap1000': '#E67E22', 'gap2000': '#6A4C93'}
    mk = {'gap0': 'o', 'gap1000': 's', 'gap2000': '^'}
    label_offsets = {'gap0': (7, -14), 'gap1000': (7, 8), 'gap2000': (8, -18)}
    fig, ax = plt.subplots(figsize=(8.8, 5.0), dpi=200)
    for r in recs:
        ax.scatter(r['a_mib_s'], r['v_mib_s'], s=16, color=cmap[r['condition']], alpha=0.30, edgecolors='none')
    for s in sums:
        c = s['condition']
        x = s['a_mib_s_median']; y = s['v_mib_s_median']
        xe = [[x - s['a_mib_s_min']], [s['a_mib_s_max'] - x]]
        ye = [[y - s['v_mib_s_min']], [s['v_mib_s_max'] - y]]
        ax.errorbar(x, y, xerr=xe, yerr=ye, fmt=mk[c], ms=11, color=cmap[c], ecolor=cmap[c], elinewidth=0.8, capsize=3, label=c)
        ax.annotate(c, (x, y), textcoords='offset points', xytext=label_offsets[c], fontproperties=fprop(9), color=cmap[c])
    ax.set_xlabel('A同窗口完成速率（MiB/s）', fontproperties=fprop(11))
    ax.set_ylabel('V速率（MiB/s）', fontproperties=fprop(11))
    ax.set_title('固定等待的双方速度取舍', fontproperties=fprop(14, True), pad=10)
    ax.set_xlim(0, 130); ax.set_ylim(0, 200)
    ax.grid(color='#DDDDDD', linewidth=0.7)
    ax.legend(prop=fprop(9), frameon=False, loc='upper left')
    caption = NL.join([
        '点：24 个原始样本（每 gap 8 个），未抖动、未连线；大符号为中位数。',
        '误差线为最小-最大范围，不是置信区间；速率单位 MiB/s。',
        'gap0/1000/2000 约等于额外等待 0/20/40 微秒（按 50 MHz 既有假设）。',
        '第 25 个样本已完成搬运与离线复核，本图每 gap 取 8 个并保留第 25 个待合并；这是阶段性描述，不是全 60、最优或显著性结论。',
    ])
    fig.text(0.02, 0.02, caption, fontproperties=fprop(8), color='#555F6D', linespacing=1.25, va='bottom')
    fig.tight_layout(rect=[0,0.18,1,1])
    fig.savefig(figs / '固定等待取舍.png', dpi=200)
    fig.savefig(figs / '固定等待取舍.pdf')
    plt.close(fig)
    return {'trials': recs, 'summary': sums, 'figure_png': str(figs / '固定等待取舍.png'), 'figure_pdf': str(figs / '固定等待取舍.pdf'), 'trials_csv': str(figs / 'p4m-trials-first24.csv'), 'summary_csv': str(figs / 'p4m-condition-summary.csv')}
ALLOWED_VISUAL = {'none', 'mechanism', 'causal', 'tradeoff', 'progress'}
def normalize_slides(obj):
    assert isinstance(obj, dict), 'slides.json must be object'
    assert obj.get('report_date') == REPORT_DATE, 'slides.report_date must be ' + REPORT_DATE
    dc = obj.get('data_cutoff')
    if isinstance(dc, dict):
        assert dc.get('board') == CUTOFF_BOARD or dc.get('latest_board') == CUTOFF_BOARD, 'data_cutoff.board mismatch'
    else:
        s = str(dc)
        assert CUTOFF_BOARD in s, 'data_cutoff string must contain ' + CUTOFF_BOARD
    slides = obj.get('slides')
    assert isinstance(slides, list) and len(slides) == EXPECTED_SLIDES, 'must have 8 slides'
    out = []
    seen = set()
    for sd in slides:
        n = int(sd.get('number'))
        assert 1 <= n <= 8 and n not in seen, 'bad slide number'
        seen.add(n)
        assert isinstance(sd.get('title'), str) and sd['title'].strip(), 'empty title slide ' + str(n)
        assert isinstance(sd.get('takeaway'), str) and sd['takeaway'].strip(), 'empty takeaway slide ' + str(n)
        b = sd.get('bullets')
        assert isinstance(b, list) and 3 <= len(b) <= 4, 'slide ' + str(n) + ' bullets must be 3-4'
        assert all(isinstance(x, str) and x.strip() for x in b), 'empty bullet slide ' + str(n)
        notes = sd.get('notes')
        if isinstance(notes, str):
            notes = [notes]
        assert isinstance(notes, list) and notes and all(isinstance(x, str) and x.strip() for x in notes), 'empty notes slide ' + str(n)
        sec = sd.get('seconds')
        assert isinstance(sec, (int, float)) and sec > 0, 'bad seconds slide ' + str(n)
        assert sd.get('visual') in ALLOWED_VISUAL, 'bad visual slide ' + str(n)
        src = sd.get('sources')
        assert isinstance(src, list) and src and all(isinstance(x, str) and x.strip() for x in src), 'bad sources slide ' + str(n)
        out.append({'number': n, 'title': sd['title'].strip(), 'takeaway': sd['takeaway'].strip(), 'bullets': [x.strip() for x in b], 'notes': [x.strip() for x in notes], 'seconds': float(sec), 'visual': sd.get('visual'), 'sources': [x.strip() for x in src]})
    out.sort(key=lambda z: z['number'])
    total = sum(s['seconds'] for s in out)
    assert MIN_SECONDS <= total <= MAX_SECONDS, 'seconds total outside 480-600: ' + str(total)
    return obj, out, total
def add_card(slide, x, y, w, h, heading, body, color=CYAN):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    sh.fill.solid(); sh.fill.fore_color.rgb = LIGHT; sh.line.color.rgb = color; sh.line.width = Pt(1.2)
    add_text(slide, x + Inches(0.15), y + Inches(0.10), w - Inches(0.3), Inches(0.42), heading, size=22, bold=True, color=color, spacing=0.98)
    body_size = 22 if (body and max(len(str(z)) for z in body) <= 28) else 20
    add_text(slide, x + Inches(0.15), y + Inches(0.62), w - Inches(0.3), h - Inches(0.74), ['• ' + str(z) for z in body], size=body_size, color=DARK_BLUE, spacing=1.05)
def build_pptx(slides_path, out_dir, e3, p4m, slides_obj, slides):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    for sd in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid(); slide.background.fill.fore_color.rgb = WHITE
        n = sd['number']
        add_title(slide, sd['title'])
        add_takeaway(slide, sd['takeaway'])
        b = sd['bullets']
        if n == 1:
            x0 = Inches(0.55); y0 = Inches(2.05); w = Inches(4.0); h = Inches(3.7)
            heads = ['已有现象', '机制证据', '连续工作量']
            for i in range(min(3, len(b))):
                add_card(slide, x0 + Inches(4.12) * i, y0, w, h, heads[i], [b[i]])
            if len(b) > 3:
                add_text(slide, Inches(0.55), Inches(5.95), Inches(12.2), Inches(0.85), '• ' + b[3], size=18, color=DARK_BLUE)
        elif n == 2:
            add_card(slide, Inches(0.55), Inches(2.05), Inches(5.95), Inches(3.8), 'E2：V 的条目被挤掉', [b[0]], ORANGE)
            add_card(slide, Inches(6.85), Inches(2.05), Inches(5.95), Inches(3.8), 'E3：记录还在，控制器被占用', [b[1]], CYAN)
            add_text(slide, Inches(0.55), Inches(5.95), Inches(12.2), Inches(0.95), ['• ' + z for z in b[2:]], size=18, color=DARK_BLUE)
        elif n == 3:
            add_card(slide, Inches(0.55), Inches(2.05), Inches(5.95), Inches(2.0), '对照：无额外延迟', ['对照轮：A 响应不额外延后。'], CYAN)
            add_card(slide, Inches(6.85), Inches(2.05), Inches(5.95), Inches(2.0), '干预：额外 128 周期', ['干预轮：只动 A 的一次页表响应时刻。'], ORANGE)
            add_text(slide, Inches(0.55), Inches(4.35), Inches(12.2), Inches(2.1), [
                '• 共用条件：64 B 诊断；V 命中条件与自身翻译不动。',
                '• 观测量：分别看 A 剩余服务与 V 剩余等待。',
                '• 限制：只控制同类目标请求；配对隐藏状态并非完全一致。',
            ], size=18, color=DARK_BLUE)
        elif n == 4:
            slide.shapes.add_picture(e3['figure_png'], Inches(0.45), Inches(2.05), width=Inches(8.65))
            add_text(slide, Inches(9.35), Inches(2.05), Inches(3.45), Inches(4.9), ['• ' + z for z in b], size=18, color=DARK_BLUE)
        elif n == 5:
            add_card(slide, Inches(0.55), Inches(2.05), Inches(4.0), Inches(3.6), 'V：1 源页＋1 目的页', ['预热 256 次', '正式 4096 × 4 KiB ＝ 16 MiB'], CYAN)
            add_card(slide, Inches(4.75), Inches(2.05), Inches(4.0), Inches(3.6), 'A：8 源页＋8 目的页', ['并行连续搬运', '完成量由硬件计数'], ORANGE)
            add_card(slide, Inches(8.95), Inches(2.05), Inches(3.85), Inches(3.6), 'gap：仅 A', ['搬完并排空后', '额外等待 0 / 20 / 40 微秒', '按 50 MHz 换算假设'], GRAY)
            add_text(slide, Inches(0.55), Inches(5.85), Inches(12.2), Inches(0.75), '• 用 V 同一测量时段统计双方进度；CPU 整批轮询，硬件连续供给。', size=18, color=DARK_BLUE)
        elif n == 6:
            slide.shapes.add_picture(p4m['figure_png'], Inches(0.35), Inches(2.05), width=Inches(7.35))
            rows = [[s['condition'].removeprefix('gap'), ('%.2f' % s['v_mib_s_median']), ('%.2f' % s['a_mib_s_median'])] for s in p4m['summary']]
            shp = slide.shapes.add_table(4, 3, Inches(8.05), Inches(2.05), Inches(4.75), Inches(2.2))
            tab = shp.table
            tab.columns[0].width = Inches(0.9)
            tab.columns[1].width = Inches(1.9)
            tab.columns[2].width = Inches(1.95)
            hdr = ['间隔', 'V中位 MiB/s', 'A中位 MiB/s']
            for j, t in enumerate(hdr):
                cell = tab.cell(0, j); cell.text = t
                for p in cell.text_frame.paragraphs:
                    for r in p.runs: set_run_font(r, 18, True, DARK_BLUE)
            for i, row in enumerate(rows, start=1):
                for j, t in enumerate(row):
                    cell = tab.cell(i, j); cell.text = str(t)
                    for p in cell.text_frame.paragraphs:
                        for r in p.runs: set_run_font(r, 18, False, DARK_BLUE)
            add_text(slide, Inches(8.05), Inches(4.5), Inches(4.75), Inches(2.2), ['• ' + z for z in b[2:]], size=18, color=DARK_BLUE)
        elif n == 7:
            bar_x = Inches(0.75); bar_y = Inches(2.25); bar_w = Inches(11.8); bar_h = Inches(0.65)
            full = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, bar_x, bar_y, bar_w, bar_h)
            full.fill.solid(); full.fill.fore_color.rgb = LIGHT; full.line.color.rgb = GRAY
            onw = int(bar_w * ONLINE / SCOPE_TOTAL); ofw = int(bar_w * OFFLINE / SCOPE_TOTAL); penw = bar_w - onw - ofw
            p1 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, bar_x, bar_y, onw, bar_h); p1.fill.solid(); p1.fill.fore_color.rgb = GREEN; p1.line.fill.background()
            p2 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, bar_x + onw, bar_y, ofw, bar_h); p2.fill.solid(); p2.fill.fore_color.rgb = ORANGE; p2.line.fill.background()
            p3 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, bar_x + onw + ofw, bar_y, penw, bar_h); p3.fill.solid(); p3.fill.fore_color.rgb = GRAY; p3.line.fill.background()
            add_text(slide, bar_x, bar_y + Inches(0.75), bar_w, Inches(0.5), '在线 24（绿） / 1 次离线复核认可（橙） / 待 35（灰）；共计划 60 次。', size=17, bold=True, color=DARK_BLUE)
            add_text(slide, Inches(0.75), Inches(4.05), Inches(11.8), Inches(2.5), ['• ' + z for z in b], size=18, color=DARK_BLUE)
        elif n == 8:
            labels = ['补数据', '决策', '独立验证']
            for i, lab in enumerate(labels):
                x = Inches(0.75 + i * 4.15)
                sh = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, x, Inches(2.25), Inches(3.55), Inches(1.35))
                sh.fill.solid(); sh.fill.fore_color.rgb = [CYAN, ORANGE, GREEN][i]; sh.line.fill.background()
                add_text(slide, x + Inches(0.25), Inches(2.45), Inches(3.05), Inches(0.95), lab, size=22, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
            add_text(slide, Inches(0.75), Inches(4.05), Inches(11.8), Inches(1.8), ['• ' + z for z in b], size=18, color=DARK_BLUE)
            add_text(slide, Inches(0.75), Inches(6.05), Inches(11.8), Inches(0.55), '老师待讨论：保护 V 到什么程度，同时要求 A 保留多少进度？', size=18, bold=True, color=ORANGE)
        else:
            raise AssertionError('unhandled slide ' + str(n))
        add_footer(slide, n, sd['sources'])
        add_notes(slide, sd)
    pptx_path = out_dir / '组会汇报.pptx'
    prs.save(str(pptx_path))
    return prs, pptx_path
def shape_bounds(prs):
    sw = int(prs.slide_width); sh = int(prs.slide_height); viol = []
    for si, slide in enumerate(prs.slides, start=1):
        for shp in slide.shapes:
            try:
                l = int(shp.left); t = int(shp.top); w = int(shp.width); h = int(shp.height)
            except Exception:
                continue
            if l < -1000 or t < -1000 or l + w > sw + 1000 or t + h > sh + 1000:
                viol.append({'slide': si, 'shape': getattr(shp, 'name', ''), 'left': l, 'top': t, 'width': w, 'height': h})
    return {'slide_width_emu': sw, 'slide_height_emu': sh, 'tolerance_emu': 1000, 'violations': viol, 'ok': len(viol) == 0}
def pptx_zip_check(path):
    try:
        z = zipfile.ZipFile(str(path))
        bad = z.testzip()
        z.close()
        return bad is None, bad
    except Exception as e:
        return False, str(e)
def reopen_check(path):
    prs = Presentation(str(path))
    notes_ok = []
    for i, slide in enumerate(prs.slides, start=1):
        txt = ''
        try:
            txt = slide.notes_slide.notes_text_frame.text
        except Exception:
            txt = ''
        notes_ok.append(bool(txt and txt.strip()))
    return {'slide_count': len(prs.slides), 'notes_nonempty': notes_ok, 'all_notes_nonempty': all(notes_ok)}
def try_pdf(out_dir, pptx_path):
    lo = shutil.which('libreoffice') or shutil.which('soffice')
    log_path = out_dir / 'convert.log'
    pdf_path = out_dir / (pptx_path.stem + '.pdf')
    if not lo:
        log_path.write_text('libreoffice not found; PDF conversion not attempted' + NL, encoding='utf-8')
        return {'attempted': False, 'ok': None, 'command': [], 'returncode': None, 'log': str(log_path), 'pdf': None, 'note': 'libreoffice missing'}
    before = pdf_path.stat().st_mtime_ns if pdf_path.exists() else None
    profile = tempfile.mkdtemp(prefix='p4m-lo-profile-', dir='/tmp')
    user_install = 'file://' + profile
    cmd = [lo, '--headless', '--norestore', '--invisible', '--nodefault', '--nofirststartwizard', '--nologo', '-env:UserInstallation=' + user_install, '--convert-to', 'pdf', '--outdir', str(out_dir), pptx_path.name]
    with log_path.open('w', encoding='utf-8') as log:
        log.write('command: ' + ' '.join(cmd) + NL)
        log.write('cwd: ' + str(out_dir) + NL)
        log.flush()
        p = subprocess.run(cmd, cwd=str(out_dir), stdout=log, stderr=subprocess.STDOUT, text=True)
    after = pdf_path.stat().st_mtime_ns if pdf_path.exists() else None
    ok = bool(p.returncode == 0 and pdf_path.exists() and pdf_path.stat().st_size > 0 and after != before)
    return {'attempted': True, 'ok': ok, 'command': cmd, 'returncode': p.returncode, 'log': str(log_path), 'pdf': str(pdf_path) if pdf_path.exists() else None, 'note': 'pdf mtime changed' if ok else 'pdf missing, empty, or not refreshed'}
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--root', default='/home/xvhaocheng/work/iommu-platform')
    ap.add_argument('--out', default=str(Path(__file__).resolve().parent))
    ap.add_argument('--slides', default=None)
    ap.add_argument('--skip-pdf', action='store_true')
    args = ap.parse_args()
    root = Path(args.root).resolve()
    out_dir = Path(args.out).resolve()
    out_dir.mkdir(parents=True, exist_ok=True)
    find_cjk_font()
    numeric_path = out_dir / NUMERIC_NAME
    if not numeric_path.exists():
        numeric_path = root / 'out/group-meeting-20261001' / NUMERIC_NAME
    slides_path = Path(args.slides).resolve() if args.slides else None
    if slides_path is None:
        for c in SLIDES_CANDIDATES:
            p = root / c
            if p.exists():
                slides_path = p; break
    if slides_path is None or not slides_path.exists():
        raise SystemExit('slides.json not found; expected one of ' + ', '.join([str(root / c) for c in SLIDES_CANDIDATES]))
    input_paths = [root / E3_RAW_REL, root / E3_REVIEW_REL, root / P4M_TRIALS_REL, root / P4M_STATS_REL, numeric_path, slides_path]
    input_paths = [p for p in input_paths if p.exists()]
    before_hashes = {rel(p, root): sha256_file(p) for p in input_paths}
    numeric = load_json(numeric_path)
    assert numeric.get('status') == 'passed', 'numeric-verification status not passed'
    for k, v in numeric.get('input_sha256', {}).items():
        p = root / k
        if p.exists():
            assert sha256_file(p) == v, 'numeric input hash mismatch ' + k
    slides_obj, slides, total_seconds = normalize_slides(load_json(slides_path))
    e3 = build_e3(root, out_dir, numeric)
    stats = load_json(root / P4M_STATS_REL)
    p4m = build_p4m(root, out_dir, stats)
    prs, pptx_path = build_pptx(slides_path, out_dir, e3, p4m, slides_obj, slides)
    bounds = shape_bounds(prs)
    assert bounds['ok'], 'shape bounding check failed: ' + str(bounds['violations'][:3])
    zip_ok, zip_bad = pptx_zip_check(pptx_path)
    reopen = reopen_check(pptx_path)
    assert reopen['slide_count'] == EXPECTED_SLIDES, 'PPTX slide count wrong'
    assert reopen['all_notes_nonempty'], 'some speaker notes empty'
    after_hashes = {rel(p, root): sha256_file(p) for p in input_paths}
    unchanged = before_hashes == after_hashes
    assert unchanged, 'input hashes changed during run'
    pdf_info = {'attempted': False, 'ok': None, 'command': [], 'returncode': None, 'log': None, 'pdf': None, 'note': 'skip-pdf requested'}
    exit_code = 0
    if not args.skip_pdf:
        pdf_info = try_pdf(out_dir, pptx_path)
        if pdf_info.get('attempted') and not pdf_info.get('ok'):
            exit_code = 1
    manifest = {'schema': 'iommu-platform.group-meeting.materials-validation.v1', 'root': str(root), 'out_dir': str(out_dir), 'report_date': REPORT_DATE, 'data_cutoff': {'latest_board': CUTOFF_BOARD, 'organized': CUTOFF_ORGANIZED}, 'input_hashes_before': before_hashes, 'input_hashes_after': after_hashes, 'inputs_unchanged': unchanged, 'slides_json': {'path': str(slides_path), 'sha256': sha256_file(slides_path), 'slide_count': len(slides), 'seconds_per_slide': [s['seconds'] for s in slides], 'seconds_total': total_seconds, 'duration_ok': MIN_SECONDS <= total_seconds <= MAX_SECONDS}, 'figures': {'e3_pairs_csv': {'path': e3['csv'], 'sha256': sha256_file(e3['csv'])}, 'mechanism_png': {'path': e3['figure_png'], 'sha256': sha256_file(e3['figure_png'])}, 'mechanism_pdf': {'path': e3['figure_pdf'], 'sha256': sha256_file(e3['figure_pdf'])}, 'p4m_trials_csv': {'path': p4m['trials_csv'], 'sha256': sha256_file(p4m['trials_csv'])}, 'p4m_summary_csv': {'path': p4m['summary_csv'], 'sha256': sha256_file(p4m['summary_csv'])}, 'tradeoff_png': {'path': p4m['figure_png'], 'sha256': sha256_file(p4m['figure_png'])}, 'tradeoff_pdf': {'path': p4m['figure_pdf'], 'sha256': sha256_file(p4m['figure_pdf'])}}, 'pptx': {'path': str(pptx_path), 'sha256': sha256_file(pptx_path), 'slide_count': reopen['slide_count'], 'notes_nonempty': reopen['notes_nonempty'], 'all_notes_nonempty': reopen['all_notes_nonempty'], 'zip_ok': zip_ok, 'zip_bad': zip_bad}, 'shape_bounding_checks': bounds, 'duration_check': {'target_seconds': TARGET_SECONDS, 'allowed_range': [MIN_SECONDS, MAX_SECONDS], 'total_seconds': total_seconds, 'ok': MIN_SECONDS <= total_seconds <= MAX_SECONDS}, 'pdf_conversion': pdf_info, 'honesty_notes': ['No board run was performed or claimed.', 'Inputs were only read; hashes before and after are recorded.', 'PDF conversion is marked true only if libreoffice actually produced a refreshed nonempty PDF.', 'The 25th P4M sample is not drawn; it is retained for later merge.', 'Figures and CSVs were computed from raw JSON, then cross-checked against review-summary and numeric-verification.']}
    write_json(out_dir / 'materials-validation.json', manifest)
    return exit_code
if __name__ == '__main__':
    sys.exit(main())
